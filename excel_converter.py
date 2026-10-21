#!/usr/bin/env python3
"""Convert a time clock export into the monthly client claim workbook.

Replaces the old VBA macros (tt_1_15 / tt_16_31 / reset).

    python excel_converter.py TIMECLOCK.xlsx --chart CLAIM_OR_CHART.xlsx
    python excel_converter.py                      # no arguments: file pickers

What it does
1. Reads the time clock export (one header row per client, shift rows below it).
2. Reads the "Family# Client# Chart" sheet (family -> clients, AHS IDs, monthly hours).
3. Writes a NEW workbook with one 32-row claim block per client, a Summary sheet,
   an Input sheet (every shift that was read) and an Issues sheet.

By default nothing is modified in place: every run builds a new file from the two input files.
With --update CLAIM.xlsx (or the "Claim workbook to update" box) a second time clock export is
written into an existing claim workbook instead; see update_workbook().
"""
from __future__ import annotations

import argparse
import calendar
import collections
import datetime as dt
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col

# --- layout of the claim sheet (matches the old workbook) -------------------
FIRST_BLOCK_ROW = 3       # row of the first block's family/client numbers
BLOCK_PITCH = 32          # rows from one block to the next
FIRST_DAY_COL = 5         # column E = day 1
LAST_DAY_COL = 35         # column AI = day 31
SUM_COL = 36              # AJ: row total
PERIOD_OFF = 18           # block offset of the 1-10/11-15/16-25/26-end subtotal header row
HALF_OFF = 24             # block offset of the 1-15/16-end subtotal header row
SUBTOTAL_LABELS = ["Homemaking", "Personal Care", "Respite", "Total"]
CLAIM_COL = 40           # AN: per-client month sum of the pink Claimed table (label; AO $$)
GRAND_COL = 37            # AK: all-category total
DEFAULT_RATE = 35.03      # $/hr, same as the old template
DEFAULT_CHART_SHEET = "Family# Client# Chart"
# config/ folder beside the script (or beside the .exe when packaged with PyInstaller)
APP_DIR = Path(sys.executable if getattr(sys, "frozen", False) else __file__).resolve().parent
if APP_DIR.parts[-3:-1] == ("excel_converter.app", "Contents"):      # macOS .app: use the folder holding the .app
    APP_DIR = APP_DIR.parents[2]
CONFIG_FILE = APP_DIR / "config" / "excel_converter.json"
OUTPUT_DIR = APP_DIR / "output"

BOLD = Font(bold=True)
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
NAME_FILL = PatternFill("solid", fgColor="FFFF00")
WARN_FONT = Font(bold=True, color="C00000")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")

# Claim sheet colors, copied from the old workbook (theme colors resolved to RGB).
YELLOW, GREY = "FFFF00", "E7E6E6"
GREEN_L, GREEN_M, GREEN_D = "E2F0D9", "C5E0B4", "A9D18E"
ORANGE_L, ORANGE_M = "FBE5D6", "F8CBAD"
BLUE_L, BLUE_M = "DEEBF7", "BDD7EE"
PINK = "FFB6C1"
# (row offset in block, first column, last column, color); later entries paint over earlier ones
BLOCK_FILLS = [
    (0, 1, 2, YELLOW),                                               # family / client
    (1, 1, 2, GREEN_L),
    (2, 1, 2, ORANGE_L), (2, 4, 4, BLUE_L),                          # client name
    (3, 1, 2, ORANGE_M), (3, 3, 4, BLUE_L),
    (4, 3, 4, BLUE_L),                                               # rate
    (7, GRAND_COL, GRAND_COL, YELLOW),                               # monthly hours total (AK)
    *[(o, 3, 4, BLUE_M) for o in range(9, 19)],                      # AHS#, claim/day, $ labels
    (10, FIRST_DAY_COL, LAST_DAY_COL, BLUE_L), (10, SUM_COL, SUM_COL, YELLOW),   # daily total
    (12, 4, 4, GREEN_D), (12, FIRST_DAY_COL, LAST_DAY_COL, BLUE_L),             # claim/day
    (12, SUM_COL, GRAND_COL, GREEN_M),
]


class InputError(Exception):
    """A problem with an input file that the user needs to fix."""


# --- data -------------------------------------------------------------------
@dataclass
class Shift:
    row: int
    start: dt.date | None
    end: dt.date | None
    hours: float


@dataclass
class ClientEntry:
    row: int
    header: str
    family: int | None
    total: float | None
    shifts: list[Shift] = field(default_factory=list)


@dataclass
class ChartRecord:
    row: int
    family: int
    client: int
    name: str
    ahs_id: object
    homemaking: float | None
    personal_care: float | None
    respite: float | None


@dataclass
class Block:
    family: int | None
    client: int | None
    name: str
    ahs_id: object
    monthly: tuple
    daily: dict          # day-of-month -> hours (only on the client that receives the hours)
    note: str = ""
    chart_name: str = ""     # name from the chart; written above the time clock name


@dataclass
class Issue:
    level: str           # "ERROR" (data is missing/wrong in the output) or "WARN"
    where: str
    message: str


# --- small helpers ----------------------------------------------------------
def to_int(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def to_float(v):
    if isinstance(v, bool) or v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return None
    return None


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
            try:
                return dt.datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def clean(text) -> str:
    """Collapse whitespace (incl. non-breaking spaces) in a name."""
    return re.sub(r"\s+", " ", str(text).replace("\xa0", " ")).strip()


def strip_leading_parens(text) -> str:
    """Drop leading (...) groups, e.g. '(回國3個月)59 Lian' -> '59 Lian'. Later parentheses stay."""
    return re.sub(r"^(?:\s*[(（][^)）]*[)）])+", "", str(text))


OFFICE = "Office"      # special header with no family number: gets a block with blank family/client/AHS


def is_office(header) -> bool:
    return clean(header).casefold() == OFFICE.casefold()


def norm(text) -> str:
    """Header text for matching: lower case, no spaces ('AHS ID #' == 'ahs id#')."""
    return re.sub(r"\s+", "", clean(text)).lower()


def find_columns(header_row, wanted: dict, where: str) -> dict:
    """Map logical names to column indexes by header text, not by fixed letters."""
    labels = {norm(v): i for i, v in enumerate(header_row) if v not in (None, "")}
    found, missing = {}, []
    for key, label in wanted.items():
        if norm(label) in labels:
            found[key] = labels[norm(label)]
        else:
            missing.append(label)
    if missing:
        raise InputError(
            f"{where}: expected column(s) {missing} but found {sorted(labels)}. "
            "Is this the right file?"
        )
    return found


# --- readers ----------------------------------------------------------------
def read_timeclock(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    issues: list[Issue] = []
    if len(wb.worksheets) > 1:
        issues.append(Issue("WARN", path.name,
                            f"Workbook has {len(wb.worksheets)} sheets; only the first "
                            f"('{wb.worksheets[0].title}') was read."))
    rows = wb.worksheets[0].iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise InputError(f"{path.name}: the sheet is empty.")
    cols = find_columns(header, {
        "name": "Type", "total": "Total shift hours",
        "start": "Start Date", "end": "End Date", "hours": "Shift hours",
    }, path.name)

    entries: list[ClientEntry] = []
    cur: ClientEntry | None = None
    for rownum, row in enumerate(rows, start=2):
        row = list(row) + [None] * (max(cols.values()) + 1 - len(row))
        name = row[cols["name"]]
        if name not in (None, ""):
            fam = re.match(r"\s*(\d+)(?=\s|\D|$)", strip_leading_parens(name))
            cur = ClientEntry(rownum, clean(name), int(fam.group(1)) if fam else None,
                              to_float(row[cols["total"]]))
            if cur.family is None and not is_office(cur.header):
                issues.append(Issue("ERROR", f"{path.name} row {rownum}",
                                    f"Header '{cur.header}' does not start with a family number; "
                                    "its hours are NOT in the claim sheet."))
            entries.append(cur)
            continue

        hours = to_float(row[cols["hours"]])
        start, end = to_date(row[cols["start"]]), to_date(row[cols["end"]])
        if hours is None and start is None:
            continue                       # blank spacer or the export's grand-total row
        if start is None and end is None:
            continue                       # total row: a number but no date (this row used to
                                           # be added to the last client by the macro)
        if cur is None:
            issues.append(Issue("ERROR", f"{path.name} row {rownum}",
                                "Shift row appears before any client header; skipped."))
            continue
        if hours is None:
            issues.append(Issue("ERROR", f"{path.name} row {rownum}",
                                f"Shift for '{cur.header}' has no hours; skipped."))
            continue
        cur.shifts.append(Shift(rownum, start, end, hours))
    wb.close()
    if not entries:
        raise InputError(f"{path.name}: no client header rows found.")
    return entries, issues


def read_chart(path: Path, sheet: str):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise InputError(f"{path.name}: no sheet named '{sheet}'. Sheets: {wb.sheetnames}. "
                         "Use --chart-sheet to name it.")
    rows = wb[sheet].iter_rows(values_only=True)
    header_idx, header = None, None
    for i, row in enumerate(rows, start=1):       # header is on row 3 today; search rows 1-10
        if any(norm(v) == "family#" for v in row if v is not None):
            header_idx, header = i, row
            break
        if i >= 10:
            break
    if header is None:
        raise InputError(f"{path.name}/{sheet}: could not find a 'Family #' header row.")
    cols = find_columns(header, {
        "family": "Family #", "client": "Client #", "name": "Name", "ahs": "AHS ID #",
        "hm": "Homemaking hours", "pc": "Personal Care Hours", "rs": "Respite hours",
    }, f"{path.name}/{sheet}")
    need = max(cols.values()) + 1
    chart: dict[int, list[ChartRecord]] = collections.defaultdict(list)
    issues: list[Issue] = []
    seen = set()
    for rownum, row in enumerate(rows, start=header_idx + 1):
        row = list(row) + [None] * (need - len(row))
        fam, cli = to_int(row[cols["family"]]), to_int(row[cols["client"]])
        if fam is None or cli is None:
            continue                       # continuation rows under each client are blank
        if (fam, cli) in seen:
            issues.append(Issue("WARN", f"{sheet} row {rownum}",
                                f"Family {fam} client {cli} appears more than once; first one used."))
            continue
        seen.add((fam, cli))
        chart[fam].append(ChartRecord(
            rownum, fam, cli, clean(row[cols["name"]] or ""), row[cols["ahs"]],
            to_float(row[cols["hm"]]), to_float(row[cols["pc"]]), to_float(row[cols["rs"]])))
    wb.close()
    if not chart:
        raise InputError(f"{path.name}/{sheet}: no family/client rows found.")
    return chart, issues


# --- planning ---------------------------------------------------------------
def pick_month(entries, date_field: str, forced: str | None):
    if forced:
        try:
            y, m = (int(p) for p in forced.split("-"))
            return y, m
        except ValueError:
            raise InputError(f"--month must look like 2026-09, got '{forced}'.")
    counts = collections.Counter()
    for e in entries:
        for s in e.shifts:
            d = s.start if date_field == "start" else s.end
            if d:
                counts[(d.year, d.month)] += 1
    if not counts:
        raise InputError("No dated shifts found, so the month can't be worked out. Use --month.")
    return counts.most_common(1)[0][0]


def plan(entries, chart, year, month, date_field):
    """Return (blocks, per-entry summary rows, input rows, issues)."""
    issues: list[Issue] = []
    by_family: dict[int, list[ClientEntry]] = collections.defaultdict(list)
    for e in entries:
        if e.family is not None:
            by_family[e.family].append(e)

    summary, input_rows = [], []
    daily_by_family: dict[int, dict[int, float]] = {}
    for e in entries:
        placed = outside = 0.0
        key = OFFICE if is_office(e.header) else e.family      # key of the claim block this entry feeds
        for s in e.shifts:
            d = s.start if date_field == "start" else s.end
            if d is None:
                issues.append(Issue("ERROR", f"time clock row {s.row}",
                                    f"{e.header}: shift has no {date_field} date; hours not placed."))
                outside += s.hours
                counted = None
            elif (d.year, d.month) != (year, month):
                outside += s.hours
                counted = None
                issues.append(Issue("WARN", f"time clock row {s.row}",
                                    f"{e.header}: {s.hours:g} h dated {d} is outside {year}-{month:02d}; "
                                    "not in the claim sheet."))
            else:
                placed += s.hours
                counted = d
                if key is not None:
                    daily_by_family.setdefault(key, {})
                    daily_by_family[key][d.day] = daily_by_family[key].get(d.day, 0) + s.hours
            input_rows.append((e.family, e.header, s.start, s.end, s.hours, s.row, counted))
        in_report = placed if key is not None else 0.0
        if key is None:
            outside = sum(s.hours for s in e.shifts)
        summary.append([e.family, e.header, e.total, len(e.shifts), round(in_report, 2),
                        round(outside, 2)])
        if e.total is not None:
            diff = round(sum(s.hours for s in e.shifts) - e.total, 2)
            if abs(diff) > 0.05:
                issues.append(Issue("WARN", f"time clock row {e.row}",
                                    f"{e.header}: shifts add to {sum(s.hours for s in e.shifts):.2f} h "
                                    f"but the time clock's Total Hours says {e.total:g}."))

    blocks: list[Block] = []
    for fam in sorted(by_family):
        group = by_family[fam]
        if len(group) > 1:
            issues.append(Issue("WARN", f"family {fam}",
                                f"{len(group)} time clock headers share this family number "
                                f"({'; '.join(g.header for g in group)}); their hours were combined."))
        name = group[0].header
        recs = chart.get(fam)
        daily = daily_by_family.get(fam, {})
        if not recs:
            issues.append(Issue("ERROR", f"family {fam}",
                                f"'{name}' is not in the chart. A block was created with no client "
                                "number, AHS ID or monthly hours. Add the family to the chart and re-run."))
            issues.append(Issue("WARN", f"family {fam}", f"'{name}': no chart name (family not in chart)."))
            blocks.append(Block(fam, None, name, None, (None, None, None), daily,
                                "NOT IN CHART - add family to chart and re-run"))
            continue
        for i, r in enumerate(recs):
            note = ""
            if len(recs) > 1:
                note = ("hours for the whole family are on this client's block" if i == 0 else
                        "second client in family - hours NOT split, enter manually")
            if not r.name:
                issues.append(Issue("WARN", f"chart row {r.row}",
                                    f"Family {fam} client {r.client} ('{name}') has no name in the chart."))
            blocks.append(Block(fam, r.client, name, r.ahs_id,
                                (r.homemaking, r.personal_care, r.respite),
                                daily if i == 0 else {}, note, r.name))
    office = [e for e in entries if is_office(e.header)]
    if office:
        if len(office) > 1:
            issues.append(Issue("WARN", OFFICE, f"{len(office)} '{OFFICE}' headers were combined."))
        blocks.append(Block(None, None, OFFICE, None, (None, None, None), daily_by_family.get(OFFICE, {})))
    return blocks, summary, input_rows, issues


# --- writer -----------------------------------------------------------------
def paint(ws, row, first_col, last_col, rgb):
    fill = PatternFill("solid", fgColor=rgb)
    for c in range(first_col, last_col + 1):
        ws.cell(row, c).fill = fill


def write_block(ws, top, b: Block, rate, days_in_month):
    r = lambda n: top + n          # noqa: E731  (row helper)
    for off, c1, c2, rgb in BLOCK_FILLS:
        paint(ws, r(off), c1, c2, rgb)
    ws.cell(top, 1, b.family)
    ws.cell(top, 2, b.client)
    ws.cell(top, 3, b.ahs_id)
    ws.cell(top, 1).font = ws.cell(top, 2).font = BOLD
    if b.note:
        c = ws.cell(top, FIRST_DAY_COL, b.note)
        c.font = WARN_FONT
        for k in (1, 2, 3):
            ws.cell(top, k).fill = WARN_FILL
    ws.cell(r(2), 2, "Client Name")
    ws.cell(r(2), 3, b.name).font = BOLD
    if b.chart_name:                                       # chart name sits above the time clock name
        ws.cell(r(1), 3, b.chart_name).font = BOLD
        ws.cell(r(1), 3).fill = NAME_FILL
    ws.cell(r(2), 4, "Days in month")
    ws.cell(r(3), 4, days_in_month)
    ws.cell(r(4), 3, "rate $/hr")
    ws.cell(r(4), 4, rate)

    labels = ["Homemaking hours", "Personal Care hours", "Respite hours"]
    for k in range(3):
        ws.cell(r(5 + k), 2, labels[k])
        if b.monthly[k] is not None:
            ws.cell(r(5 + k), 4, b.monthly[k])
    ws.cell(r(8), 2, f"=SUM(D{r(5)}:D{r(7)})")
    ws.cell(r(8), 4, "Adj. hrs")
    ws.cell(r(9), 2, "AHS#")
    ws.cell(r(9), 3, f"=C{top}")
    ws.cell(r(10), 4, "Total")
    ws.cell(r(12), 4, "Claim/day")
    for k, lab in enumerate(["Homemaking $$", "Personal Care $$", "Respite $$", "Total $$ "]):
        ws.cell(r(14 + k), 4, lab)

    for day, hrs in b.daily.items():                       # personal care row, as the macro did
        c = ws.cell(r(6), FIRST_DAY_COL + day - 1, round(hrs, 2))
        c.number_format = "0.00"
    rate_ref = f"$D{r(4)}"
    for c in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
        L = col(c)
        ws.cell(r(10), c, f"=SUM({L}{r(5)}:{L}{r(8)})")
        ws.cell(r(12), c, f"={L}{r(10)}*{rate_ref}")
        for k in range(3):
            ws.cell(r(14 + k), c, f"={L}{r(5 + k)}*{rate_ref}")
        ws.cell(r(17), c, f"=SUM({L}{r(14)}:{L}{r(16)})")
        for rr in (10, 12, 14, 15, 16, 17):
            ws.cell(r(rr), c).number_format = "0.00"
    first, last = col(FIRST_DAY_COL), col(LAST_DAY_COL)
    for rr in (5, 6, 7, 8, 10, 12):
        ws.cell(r(rr), SUM_COL, f"=SUM({first}{r(rr)}:{last}{r(rr)})").number_format = "0.00"
    for rr in (14, 15, 16, 17):                            # dollar rows: month total after the last day
        ws.cell(r(rr), SUM_COL, f"=SUM({first}{r(rr)}:{last}{r(rr)})").number_format = "0.00"
    ws.cell(r(7), GRAND_COL, f"=SUM({col(SUM_COL)}{r(5)}:{col(SUM_COL)}{r(7)})")
    ws.cell(r(12), GRAND_COL, f"={col(SUM_COL)}{r(10)}*{rate_ref}").number_format = "0.00"

    # subtotal tables: 10/5-day periods (header +18, rows +19..+22), then half months (header +24, rows +25..+28).
    # Per span: label at dest-2, hours at dest-1, dollars at dest; month sums sit right of the last day column.
    def block_formula(k, dollars, d1, d2, c):
        row = r(14 + k) if dollars else r(5 + k)
        return f"=SUM({col(FIRST_DAY_COL + d1 - 1)}{row}:{col(FIRST_DAY_COL + d2 - 1)}{row})"
    for off, spans in ((PERIOD_OFF, period_spans(days_in_month)), (HALF_OFF, half_spans(days_in_month))):
        write_subtotals(ws, r(off), spans, block_formula, "Days ", BLUE_M, "Month ",
                        claimed=(lambda k, c: None) if off == HALF_OFF else None)


def period_spans(days):
    return [(1, 10), (11, 15), (16, 25), (26, days)]


def half_spans(days):
    return [(1, 15), (16, days)]


def claimed_cols(dest):
    """Label and $$ columns of the Claimed box for a span ending in column dest.
    The box covers three columns (label spills into the middle one); dest-4, dest-3 stay empty."""
    return dest - 7, dest - 5


def write_claimed(ws, header_row, spans, claimed, grand=False):
    """Pink box per span: Claimed 1-3 (blank, typed by hand) and Total Claimed (their sum, shown right).
    For the all-clients table (`grand`) only the yellow Total Claimed row is drawn: the sum of each
    client's Total Claimed."""
    right = Alignment(horizontal="right")
    fill = PatternFill("solid", fgColor=YELLOW if grand else PINK)
    dests = [FIRST_DAY_COL + d2 - 1 for _, d2 in spans]
    for dest in dests:
        lab_c, d_c = claimed_cols(dest)
        first_k = 3 if grand else 0
        top = header_row + 1 + first_k
        for k in range(first_k, 4):
            row = header_row + 1 + k
            ws.cell(row, lab_c, f"Claimed {k + 1}" if k < 3 else "Total Claimed")
            for c in range(lab_c, d_c + 1):
                ws.cell(row, c).fill = fill
            cell = ws.cell(row, d_c)
            if k == 3:
                cell.value = claimed(k, d_c) if grand else f"=SUM({col(d_c)}{header_row + 1}:{col(d_c)}{header_row + 3})"
                ws.cell(row, lab_c).font = cell.font = BOLD if grand else Font()
            cell.number_format = "0.00"
        box(ws, top, header_row + 4, lab_c, d_c)
    ws.cell(header_row, CLAIM_COL + 1, "Claimed $$").font = BOLD
    ws.cell(header_row, CLAIM_COL + 1).alignment = right
    k = 3
    row = header_row + 1 + k
    ws.cell(row, CLAIM_COL, "Month Total Claimed").alignment = right
    c = ws.cell(row, CLAIM_COL + 1, "=" + "+".join(f"{col(claimed_cols(d)[1])}{row}" for d in dests))
    c.number_format = "0.00"
    c.font = BOLD
    ws.cell(row, CLAIM_COL).fill = c.fill = fill


def write_subtotals(ws, header_row, spans, formula, title_prefix, label_rgb, month_prefix, claimed=None):
    """One subtotal table: a column group per span plus the month sums (AJ label, AK hours, AL $$).
    `formula(k, dollars, d1, d2, c)` gives the cell formula for category k (0-2) in column c.
    `claimed(k, c)`, when given, adds a pink "Claimed" table left of each span (label, $$) with a
    two-cell gap before the blue subtotal labels, plus its month sums right of the Month columns."""
    right = Alignment(horizontal="right")
    ws.cell(header_row, SUM_COL, "Month").font = BOLD
    ws.cell(header_row, SUM_COL).alignment = right
    ws.cell(header_row, GRAND_COL, "Hours").font = BOLD
    ws.cell(header_row, GRAND_COL + 1, "$$").font = BOLD
    dests = []
    for d1, d2 in spans:
        dest = FIRST_DAY_COL + d2 - 1
        dests.append(dest)
        ws.cell(header_row, dest - 2, f"{title_prefix}{d1}-{d2}").font = BOLD
        ws.cell(header_row, dest - 2).alignment = right
        ws.cell(header_row, dest - 1, "Hours").font = BOLD
        ws.cell(header_row, dest, "$$").font = BOLD
        for k, lab in enumerate(SUBTOTAL_LABELS):
            row = header_row + 1 + k
            ws.cell(row, dest - 2, f"Subtotal {lab}").alignment = right   # right-aligned so the text spills left
            paint(ws, row, dest - 2, dest - 2, label_rgb)
            for c in (dest - 1, dest):
                if k < 3:
                    ws.cell(row, c, formula(k, c == dest, d1, d2, c))
                else:
                    L = col(c)
                    ws.cell(row, c, f"=SUM({L}{header_row + 1}:{L}{header_row + 3})").font = BOLD
                ws.cell(row, c).number_format = "0.00"
    if claimed:
        write_claimed(ws, header_row, spans, claimed, grand=title_prefix.startswith("All"))
    for k, lab in enumerate(SUBTOTAL_LABELS):
        row = header_row + 1 + k
        ws.cell(row, SUM_COL, f"{month_prefix}{lab}").alignment = right
        paint(ws, row, SUM_COL, SUM_COL, YELLOW)
        for out, shift in ((GRAND_COL, 1), (GRAND_COL + 1, 0)):          # hours, then $$
            ws.cell(row, out, "=" + "+".join(f"{col(d - shift)}{row}" for d in dests)).number_format = "0.00"
            if k == 3:
                ws.cell(row, out).font = BOLD


def box(ws, top_row, bottom_row, first_col, last_col):
    """Outline a range with a medium border, keeping each cell's fill and font."""
    edge = Side(style="medium")
    for row in range(top_row, bottom_row + 1):
        for c in range(first_col, last_col + 1):
            ws.cell(row, c).border = Border(left=edge if c == first_col else None,
                                            right=edge if c == last_col else None,
                                            top=edge if row == top_row else None,
                                            bottom=edge if row == bottom_row else None)


def write_grand_totals(ws, tops, days_in_month):
    """Below the last block: each subtotal table (hours and $$) summed over all clients."""
    last_top = tops[-1]
    header = last_top + HALF_OFF + 6                       # one blank row under the last block's box
    for off, spans in ((PERIOD_OFF, period_spans(days_in_month)), (HALF_OFF, half_spans(days_in_month))):
        def all_clients(k, dollars, d1, d2, c, off=off):
            refs = ",".join(f"{col(c)}{t + off + 1 + k}" for t in tops)
            return f"=SUM({refs})"
        def all_claimed(k, c, off=off):
            refs = ",".join(f"{col(c)}{t + off + 1 + k}" for t in tops)
            return f"=SUM({refs})"
        write_subtotals(ws, header, spans, all_clients, "All clients ", YELLOW, "Month ",
                        claimed=all_claimed if off == HALF_OFF else None)
        header += 6


HEAD = {"Summary": (["Family", "Time clock header", "Time Clock 'Total Hours'", "Shift rows",
                     "Hours in claim sheet", "Hours NOT in claim sheet"], (9, 48, 20, 11, 20, 24)),
        "Input": (["Family", "Time clock header", "Start Date", "End Date", "Shift hours",
                   "Time clock row", "Date used"], (9, 48, 12, 12, 12, 14, 12)),
        "Issues": (["Level", "Where", "Message"], (8, 26, 120))}


def get_sheet(wb, title):
    """The named sheet, created with its header row and column widths if the workbook lacks it."""
    if title in wb.sheetnames:
        return wb[title]
    header, widths = HEAD[title]
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD_FILL
    for letter, w in zip("ABCDEFG", widths):
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A2"
    return ws


def add_summary(wb, summary):
    """Append summary rows, keeping a single TOTAL row at the bottom."""
    s = get_sheet(wb, "Summary")
    for r_ in range(s.max_row, 1, -1):
        if s.cell(r_, 1).value == "TOTAL":
            s.delete_rows(r_)
    for row in summary:
        s.append(row)
    n = s.max_row
    s.append(["TOTAL", None, f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", f"=SUM(E2:E{n})", f"=SUM(F2:F{n})"])
    for c in s[n + 1]:
        c.font = BOLD


def add_input(wb, input_rows):
    i = get_sheet(wb, "Input")
    first = i.max_row + 1
    for row in sorted(input_rows, key=lambda x: (x[0] is None, x[0] or 0, x[5])):
        i.append(list(row))
    for col_idx in (3, 4, 7):
        for r_ in i.iter_rows(min_row=first, min_col=col_idx, max_col=col_idx):
            for c in r_:
                c.number_format = "yyyy-mm-dd"


def add_issues(wb, issues):
    """Append issues below the ones already on the sheet (never replaces them)."""
    x = get_sheet(wb, "Issues")
    if x.max_row == 2 and x.cell(2, 1).value == "OK":
        x.delete_rows(2)
    for iss in sorted(issues, key=lambda z: (z.level != "ERROR", z.level != "WARN", z.where)):
        x.append([iss.level, iss.where, iss.message])
        if iss.level == "ERROR":
            x.cell(x.max_row, 1).font = WARN_FONT
    if x.max_row == 1:
        x.append(["OK", "", "No issues found."])


def write_workbook(out_path, blocks, summary, input_rows, issues, year, month, rate, date_field):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Claim {year}-{month:02d}"
    days = calendar.monthrange(year, month)[1]

    ws["A2"], ws["B2"] = "Family Number", "Client Number"
    ws["D2"] = rate
    ws.cell(2, FIRST_DAY_COL, dt.date(year, month, 1)).number_format = "m/d"
    for c in range(FIRST_DAY_COL + 1, LAST_DAY_COL + 1):
        prev = col(c - 1)
        ws.cell(2, c, f'=IFERROR(IF({prev}2=EOMONTH({prev}2,0),"",{prev}2+1),"")').number_format = "m/d"
    for c in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
        ws.cell(1, c, f'=TEXT({col(c)}2,"ddd")').alignment = Alignment(horizontal="center")
    for c in ("A2", "B2", "D2"):
        ws[c].font = BOLD
    paint(ws, 1, FIRST_DAY_COL, LAST_DAY_COL, GREY)
    paint(ws, 2, 4, LAST_DAY_COL, GREEN_L)
    ws.cell(2, FIRST_DAY_COL).fill = PatternFill("solid", fgColor=YELLOW)
    ws["A1"] = f"Hours counted by shift {date_field} date. Generated by excel_converter.py."
    ws["A1"].font = Font(italic=True, color="666666")

    tops = [FIRST_BLOCK_ROW + i * BLOCK_PITCH for i in range(len(blocks))]
    for top, b in zip(tops, blocks):
        write_block(ws, top, b, rate, days)
        box(ws, top, top + HALF_OFF + 4, 1, CLAIM_COL + 1)
    if blocks:
        write_grand_totals(ws, tops, days)
    for c, w in zip("ABCD", (13, 14, 30, 14)):
        ws.column_dimensions[c].width = w
    for c in range(FIRST_DAY_COL, GRAND_COL + 2):
        ws.column_dimensions[col(c)].width = 7.5
    ws.column_dimensions[col(CLAIM_COL)].width = 24
    ws.column_dimensions[col(SUM_COL)].width = 24              # AJ also holds the "Month ..." labels
    ws.freeze_panes = "A3"

    add_summary(wb, summary)
    add_input(wb, input_rows)
    add_issues(wb, issues)
    wb.save(out_path)


def block_key(family, client):
    return ("office",) if family is None and client is None else (family, client)


def month_of(ws):
    d = ws.cell(2, FIRST_DAY_COL).value
    if not isinstance(d, (dt.date, dt.datetime)):
        raise InputError(f"{ws.title}: cell E2 should hold the first day of the month.")
    return d.year, d.month


def open_claim_workbook(path: Path):
    """Load an existing claim workbook; return (workbook, claim sheet, year, month, block top rows)."""
    if path.suffix.lower() != ".xlsx":
        raise InputError(f"{path.name}: only .xlsx claim workbooks can be updated.")
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:                                   # openpyxl raises several types for bad files
        raise InputError(f"{path.name}: cannot open as a workbook ({e}).") from e
    sheets = [n for n in wb.sheetnames if n.startswith("Claim ")]
    if not sheets:
        raise InputError(f"{path.name}: no 'Claim YYYY-MM' sheet. Choose a workbook made by this program.")
    ws = wb[sheets[0]]
    year, month = month_of(ws)
    tops = [r - 2 for r in range(FIRST_BLOCK_ROW + 2, ws.max_row + 1) if ws.cell(r, 2).value == "Client Name"]
    if not tops:
        raise InputError(f"{path.name}/{ws.title}: no client blocks found.")
    if ws.cell(tops[0] + HALF_OFF, FIRST_DAY_COL + 12).value != "Days 1-15":
        raise InputError(f"{path.name}/{ws.title}: blocks do not match the current layout "
                         "(made by an older version?). Build a new workbook instead.")
    return wb, ws, year, month, tops


def update_workbook(wb, ws, tops, blocks, summary, input_rows, issues, year, month, rate):
    """Write the planned blocks into an open claim workbook. New time clock hours replace what is there
    (a warning for each non-zero value replaced); clients not in the file get new blocks."""
    days = calendar.monthrange(year, month)[1]
    existing: dict[tuple, list[int]] = collections.defaultdict(list)
    for top in tops:
        existing[block_key(to_int(ws.cell(top, 1).value), to_int(ws.cell(top, 2).value))].append(top)

    new_blocks = []
    for b in blocks:
        key = block_key(b.family, b.client)
        rows = existing.get(key)
        if not rows:
            new_blocks.append(b)
            continue
        top = rows[0]
        if len(rows) > 1:
            issues.append(Issue("WARN", f"{ws.title} row {top}",
                                f"{b.name}: the file has {len(rows)} blocks for this client (rows "
                                f"{', '.join(str(t) for t in rows)}); only the first was written to."))
        for day, hrs in sorted(b.daily.items()):
            cell = ws.cell(top + 6, FIRST_DAY_COL + day - 1)
            new, old = round(hrs, 2), cell.value
            blank = old is None or (isinstance(old, str) and not old.strip())
            zero = isinstance(old, (int, float)) and abs(old) < 0.005
            same = isinstance(old, (int, float)) and abs(old - new) < 0.005
            if not (blank or zero or same):
                issues.append(Issue("WARN", f"{ws.title}!{cell.coordinate}",
                                    f"{b.name}: day {day} had {old} h; replaced with {new:g} h from the new time clock."))
            cell.value = new
            cell.number_format = "0.00"

    if new_blocks:
        file_rate = ws.cell(2, 4).value
        if isinstance(file_rate, (int, float)) and abs(file_rate - rate) > 1e-9:
            issues.append(Issue("WARN", f"{ws.title}!D2",
                                f"New blocks use the file's rate {file_rate:g}, not {rate:g}."))
            rate = file_rate
        start = tops[-1] + HALF_OFF + 5                      # first row under the last block's box
        if start <= ws.max_row:
            ws.delete_rows(start, ws.max_row - start + 1)     # old all-clients totals; rebuilt below
        for b in new_blocks:
            top = tops[-1] + BLOCK_PITCH
            tops.append(top)
            write_block(ws, top, b, rate, days)
            box(ws, top, top + HALF_OFF + 4, 1, CLAIM_COL + 1)
            issues.append(Issue("INFO", f"{ws.title} row {top}", f"{b.name}: not in the file, new block added."))
        write_grand_totals(ws, tops, days)

    add_summary(wb, summary)
    add_input(wb, input_rows)
    add_issues(wb, issues)
    return len(new_blocks)


# --- command line / GUI -----------------------------------------------------
def default_output() -> Path:
    OUTPUT_DIR.mkdir(exist_ok=True)
    return OUTPUT_DIR / f"{dt.datetime.now():%Y-%m-%d_%H%M%S}.xlsx"


def run(timeclock: Path, chart_path: Path, chart_sheet: str, out: Path | None,
        month: str | None, rate: float, date_field: str, update: Path | None = None):
    """Build a claim workbook from the time clock. With `update`, write into that existing claim workbook
    instead. The chosen file is left as it is; the result goes to `out` (default: a new file in output/)."""
    entries, issues = read_timeclock(timeclock)
    chart, chart_issues = read_chart(chart_path, chart_sheet)
    issues += chart_issues
    if update:
        wb, ws, year, mon, tops = open_claim_workbook(update)
        if month and month != f"{year}-{mon:02d}":
            raise InputError(f"--month {month} does not match the workbook's month {year}-{mon:02d}.")
    else:
        year, mon = pick_month(entries, date_field, month)
    blocks, summary, input_rows, plan_issues = plan(entries, chart, year, mon, date_field)
    issues += plan_issues
    new_blocks = 0
    if out is None:
        out = default_output()
    try:
        if update:
            new_blocks = update_workbook(wb, ws, tops, blocks, summary, input_rows, issues, year, mon, rate)
            wb.save(out)                                 # the chosen file itself is never modified
        else:
            write_workbook(out, blocks, summary, input_rows, issues, year, mon, rate, date_field)
    except PermissionError as e:
        raise InputError(f"Cannot save {out}. Close it in Excel and try again.") from e
    total_in = sum(r[4] for r in summary)
    total_out = sum(r[5] for r in summary)
    return out, dict(blocks=len(blocks), clients=len(entries), month=f"{year}-{mon:02d}",
                     hours_in=round(total_in, 2), hours_out=round(total_out, 2),
                     updated=bool(update), new_blocks=new_blocks,
                     errors=sum(i.level == "ERROR" for i in issues),
                     warnings=sum(i.level == "WARN" for i in issues))


def read_config() -> dict:
    """Contents of config/excel_converter.json, or {} if missing or unreadable."""
    try:
        cfg = json.loads(CONFIG_FILE.read_text())
    except (OSError, ValueError):
        return {}
    return cfg if isinstance(cfg, dict) else {}


def config_chart() -> str | None:
    """Chart path saved in config/excel_converter.json, if it is set and the file exists."""
    chart = read_config().get("chart")
    return chart if chart and Path(chart).exists() else None


def config_rate() -> float:
    """default_rate from config/excel_converter.json, else DEFAULT_RATE."""
    rate = read_config().get("default_rate")
    if rate is None:
        return DEFAULT_RATE
    if isinstance(rate, bool) or not isinstance(rate, (int, float)) or rate <= 0:
        raise InputError(f'"default_rate" in {CONFIG_FILE} must be a positive number, got {rate!r}')
    return float(rate)


def gui_pick(args):
    """Show one form for the time clock, chart, rate and output; fill args from it."""
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
    except ImportError:
        sys.exit("No window available here. Run: python excel_converter.py TIMECLOCK.xlsx --chart FILE.xlsx")
    ft = [("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")]
    saved_chart = config_chart()
    try:
        default_rate = config_rate()
    except InputError:
        default_rate = DEFAULT_RATE

    root = tk.Tk()
    root.title("Excel converter")
    root.columnconfigure(1, weight=1)
    tc_var = tk.StringVar()
    chart_var = tk.StringVar(value=args.chart or saved_chart or "")
    rate_var = tk.StringVar(value=f"{args.rate if args.rate is not None else default_rate:g}")
    out_var = tk.StringVar(value=args.out or "")
    upd_var = tk.StringVar(value=args.update or "")
    result = {}

    def browse(var, title, save=False):
        if save:
            path = filedialog.asksaveasfilename(title=title, defaultextension=".xlsx",
                                                filetypes=[("Excel files", "*.xlsx")])
        else:
            path = filedialog.askopenfilename(title=title, filetypes=ft)
        if path:
            var.set(path)

    rows = [("Time Clock", tc_var, "Time Clock", False, "Required"),
            ("Client chart workbook", chart_var, "Workbook containing the Family# Client# Chart sheet", False,
             "Default from config" if saved_chart else "Required"),
            ("Claim workbook to update", upd_var, "Existing claim workbook to copy and add to", False,
             "Leave blank to create a new file"),
            ("Output file", out_var, "Save the claim workbook as", True,
             "Leave blank for output/<date>_<time>.xlsx (the workbook above is never changed)")]
    entries = {}
    for i, (label, var, title, save, hint) in enumerate(rows):
        ttk.Label(root, text=label).grid(row=i * 2, column=0, sticky="w", padx=8, pady=(8, 0))
        entries[label] = ttk.Entry(root, textvariable=var, width=60)
        entries[label].grid(row=i * 2, column=1, sticky="ew", padx=4, pady=(8, 0))
        ttk.Button(root, text="Browse...",
                   command=lambda v=var, t=title, s=save: browse(v, t, s)).grid(row=i * 2, column=2, padx=8, pady=(8, 0))
        ttk.Label(root, text=hint, foreground="gray").grid(row=i * 2 + 1, column=1, sticky="w", padx=4)
    r = len(rows) * 2
    ttk.Label(root, text="Rate ($/hr)").grid(row=r, column=0, sticky="w", padx=8, pady=8)
    ttk.Entry(root, textvariable=rate_var, width=10).grid(row=r, column=1, sticky="w", padx=4, pady=8)

    def ok(_event=None):
        tc, chart, out = tc_var.get().strip(), chart_var.get().strip(), out_var.get().strip()
        upd = upd_var.get().strip()
        if upd and not Path(upd).is_file():
            return messagebox.showerror("Missing file", "The claim workbook to update was not found.", parent=root)
        if not tc or not Path(tc).is_file():
            return messagebox.showerror("Missing file", "Choose the Time Clock file.", parent=root)
        if not chart or not Path(chart).is_file():
            return messagebox.showerror("Missing file", "Choose the client chart workbook.", parent=root)
        try:
            rate = float(rate_var.get())
            if rate <= 0 or rate != rate:
                raise ValueError
        except ValueError:
            return messagebox.showerror("Bad rate", "The rate must be a positive number.", parent=root)
        result.update(timeclock=tc, chart=chart, rate=rate, out=out or None, update=upd or None)
        root.withdraw()
        root.quit()

    buttons = ttk.Frame(root)
    buttons.grid(row=r + 1, column=0, columnspan=3, sticky="e", padx=8, pady=(0, 8))
    ttk.Button(buttons, text="Cancel", command=root.quit).pack(side="right", padx=(8, 0))
    ttk.Button(buttons, text="Convert", command=ok).pack(side="right")
    root.bind("<Return>", ok)
    root.bind("<Escape>", lambda _e: root.quit())
    root.protocol("WM_DELETE_WINDOW", root.quit)
    root.mainloop()
    if not result:
        root.destroy()
        sys.exit(0)

    if result["chart"] != saved_chart:
        try:
            CONFIG_FILE.parent.mkdir(exist_ok=True)
            CONFIG_FILE.write_text(json.dumps({"default_rate": DEFAULT_RATE, **read_config(),
                                               "chart": result["chart"]}, indent=2))
        except OSError:
            pass
    args.timeclock, args.chart, args.rate, args.out, args.update = (
        result[k] for k in ("timeclock", "chart", "rate", "out", "update"))
    return root, messagebox


def main(argv=None):
    p = argparse.ArgumentParser(description="Build the client claim workbook from a time clock file.")
    p.add_argument("timeclock", nargs="?", help="time clock file (.xlsx)")
    p.add_argument("--chart", help="workbook with the 'Family# Client# Chart' sheet (.xlsx/.xlsm); "
                                   "default: the \"chart\" path in config/excel_converter.json")
    p.add_argument("--chart-sheet", default=DEFAULT_CHART_SHEET)
    p.add_argument("-o", "--out", help="output file (default: output/<date>_<time>.xlsx beside the script)")
    p.add_argument("--update", help="existing claim workbook (made by this program) to write into; "
                                    "default: build a new file. The original is left untouched; the result goes to output/ (or -o).")
    p.add_argument("--month", help="claim month as YYYY-MM (default: the month most shifts fall in)")
    p.add_argument("--rate", type=float, default=None,
                   help=f"$/hr (default: \"default_rate\" in config/excel_converter.json, else {DEFAULT_RATE})")
    p.add_argument("--date-field", choices=("start", "end"), default="start",
                   help="which date a shift counts on; matters for overnight shifts (default: start)")
    args = p.parse_args(argv)

    gui = None
    if not args.timeclock:
        gui = gui_pick(args)
    elif not args.chart:
        args.chart = config_chart()
        if not args.chart:
            p.error(f"--chart is required (or set \"chart\" to an existing file in {CONFIG_FILE})")
    try:
        rate = args.rate if args.rate is not None else config_rate()
        out, stats = run(Path(args.timeclock), Path(args.chart), args.chart_sheet,
                         Path(args.out) if args.out else None, args.month, rate, args.date_field,
                         Path(args.update) if args.update else None)
    except InputError as e:
        if gui:
            gui[1].showerror("Cannot convert", str(e))
        sys.exit(f"ERROR: {e}")
    msg = f"{'Updated' if stats['updated'] else 'Wrote'} {out}\n"
    if stats["updated"]:
        msg += f"{stats['new_blocks']} new client block(s) added.\n"
    msg += (f"{stats['blocks']} claim blocks for {stats['clients']} time clock clients "
            f"({stats['month']}); {stats['hours_in']:g} h in the claim sheet, "
            f"{stats['hours_out']:g} h left out.\n"
            f"{stats['errors']} error(s), {stats['warnings']} warning(s) - see the Issues sheet.")
    print(msg)
    if gui:
        gui[1].showinfo("Done", msg)
    return 1 if stats["errors"] else 0


if __name__ == "__main__":
    sys.exit(main())
