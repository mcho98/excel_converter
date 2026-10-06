"""Tests build tiny synthetic workbooks, so no real client data is needed."""
import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import ct_claims as ct  # noqa: E402

D = dt.datetime


def make_timeclock(path, blocks, total_row=True):
    """blocks: [(header, total, [(start, end, hours), ...])]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Client Number and Name", None, None, "Total Hours", "First name", "Last name",
               "Start Date", "In", None, "End Date", "Out", None, "Shift hours", None])
    for header, total, shifts in blocks:
        ws.append([header, None, None, total] + [""] * 9)
        for s, e, h in shifts:
            ws.append(["", None, None, "", "A", "B", s, "09:00", None, e, "15:00", None, h])
    if total_row:
        ws["M%d" % (ws.max_row + 1)] = "=SUM(M2:M%d)" % ws.max_row
    wb.save(path)


def make_chart(path, rows):
    """rows: [(family, client, name, ahs, hm, pc, rs)]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ct.DEFAULT_CHART_SHEET
    ws["A2"] = "Code"
    ws.append([])
    ws.cell(3, 1, "Family #"), ws.cell(3, 2, "Client #"), ws.cell(3, 4, "Name")
    ws.cell(3, 6, "AHS ID #"), ws.cell(3, 10, "Homemaking hours")
    ws.cell(3, 11, "Personal Care Hours"), ws.cell(3, 12, "Respite hours")
    r = 4
    for fam, cli, name, ahs, hm, pc, rs in rows:
        for c, v in ((1, fam), (2, cli), (4, name), (6, ahs), (10, hm), (11, pc), (12, rs)):
            ws.cell(r, c, v)
        r += 4                        # real chart leaves 3 blank rows under each client
    wb.save(path)


@pytest.fixture
def files(tmp_path):
    tc, ch = tmp_path / "tc.xlsx", tmp_path / "chart.xlsx"
    make_timeclock(tc, [
        ("5  Bob Smith", 14, [(D(2026, 9, 17), D(2026, 9, 17), 6), (D(2026, 9, 20), D(2026, 9, 21), 8)]),
        ("2 Al Jones", 4, [(D(2026, 9, 18), D(2026, 9, 18), 4)]),
        ("99 New Person", 3, [(D(2026, 9, 19), D(2026, 9, 19), 3)]),
    ])
    make_chart(ch, [(2, 3, "Al", "A1", 1, 2, 3), (5, 6, "Bob", "B1", 0, 5, 1), (5, 7, "Bobs wife", "B2", 0, 5, 1)])
    return tc, ch, tmp_path / "out.xlsx"


def run(files, **kw):
    tc, ch, out = files
    args = dict(chart_sheet=ct.DEFAULT_CHART_SHEET, out=out, month=None, rate=34.03, date_field="start")
    args.update(kw)
    return ct.run(tc, ch, args.pop("chart_sheet"), args.pop("out"), args.pop("month"),
                  args.pop("rate"), args.pop("date_field"))


def sheet(out):
    return openpyxl.load_workbook(out)


def test_total_row_is_not_counted(files):
    _, stats = run(files)
    assert stats["hours_in"] == 21 and stats["hours_out"] == 0     # 14 + 4 + 3, not + 21 again


def test_blocks_sorted_and_hours_on_start_date(files):
    out, _ = run(files)
    ws = sheet(out).worksheets[0]
    fams = [ws.cell(ct.FIRST_BLOCK_ROW + i * ct.BLOCK_PITCH, 1).value for i in range(5)]
    assert fams == [2, 5, 5, 99, None]                  # sorted by family; family 5 has two clients
    first5 = ct.FIRST_BLOCK_ROW + ct.BLOCK_PITCH
    pc = first5 + 6
    assert ws.cell(pc, ct.FIRST_DAY_COL + 16).value == 6            # 17th
    assert ws.cell(pc, ct.FIRST_DAY_COL + 19).value == 8            # overnight shift stays on the 20th
    assert ws.cell(first5 + ct.BLOCK_PITCH + 6, ct.FIRST_DAY_COL + 16).value is None   # 2nd client empty


def test_end_date_option_moves_overnight_shift(files):
    out, _ = run(files, date_field="end")
    ws = sheet(out).worksheets[0]
    pc = ct.FIRST_BLOCK_ROW + ct.BLOCK_PITCH + 6
    assert ws.cell(pc, ct.FIRST_DAY_COL + 20).value == 8            # 21st


def test_family_missing_from_chart_is_flagged_not_dropped(files):
    out, stats = run(files)
    assert stats["errors"] == 1
    wb = sheet(out)
    msgs = [r[2] for r in wb["Issues"].iter_rows(min_row=2, values_only=True)]
    assert any("99" in m and "not in the chart" in m for m in msgs)
    ws = wb.worksheets[0]
    top = ct.FIRST_BLOCK_ROW + 3 * ct.BLOCK_PITCH                   # family 99 block
    assert ws.cell(top, 1).value == 99 and ws.cell(top + 6, ct.FIRST_DAY_COL + 18).value == 3


def test_shifts_outside_month_are_reported(tmp_path):
    tc, ch = tmp_path / "tc.xlsx", tmp_path / "chart.xlsx"
    make_timeclock(tc, [("2 Al", 10, [(D(2026, 9, 30), D(2026, 10, 1), 6), (D(2026, 10, 1), D(2026, 10, 1), 4)])])
    make_chart(ch, [(2, 3, "Al", "A1", 1, 2, 3)])
    _, stats = ct.run(tc, ch, ct.DEFAULT_CHART_SHEET, tmp_path / "o.xlsx", "2026-09", 34.03, "start")
    assert stats["hours_in"] == 6 and stats["hours_out"] == 4 and stats["warnings"] >= 1


def test_same_family_twice_is_combined(tmp_path):
    tc, ch = tmp_path / "tc.xlsx", tmp_path / "chart.xlsx"
    make_timeclock(tc, [("2 Al", 2, [(D(2026, 9, 1), D(2026, 9, 1), 2)]),
                        ("2 Al again", 3, [(D(2026, 9, 1), D(2026, 9, 1), 3)])])
    make_chart(ch, [(2, 3, "Al", "A1", 1, 2, 3)])
    out = tmp_path / "o.xlsx"
    ct.run(tc, ch, ct.DEFAULT_CHART_SHEET, out, None, 34.03, "start")
    ws = sheet(out).worksheets[0]
    assert ws.cell(ct.FIRST_BLOCK_ROW + 6, ct.FIRST_DAY_COL).value == 5


def test_header_without_family_number_is_an_error(tmp_path):
    tc, ch = tmp_path / "tc.xlsx", tmp_path / "chart.xlsx"
    make_timeclock(tc, [("Office", 2, [(D(2026, 9, 1), D(2026, 9, 1), 2)])])
    make_chart(ch, [(2, 3, "Al", "A1", 1, 2, 3)])
    _, stats = ct.run(tc, ch, ct.DEFAULT_CHART_SHEET, tmp_path / "o.xlsx", "2026-09", 34.03, "start")
    assert stats["errors"] == 1 and stats["hours_out"] == 2


def test_wrong_file_gives_readable_error(tmp_path):
    bad = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Qty"])
    wb.save(bad)
    with pytest.raises(ct.InputError, match="Shift hours"):
        ct.read_timeclock(bad)
