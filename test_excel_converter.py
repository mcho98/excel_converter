"""Tests build tiny synthetic workbooks, so no real client data is needed."""
import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import excel_converter as ct  # noqa: E402

D = dt.datetime


def make_timeclock(path, blocks, total_row=True):
    """blocks: [(header, total, [(start, end, hours), ...])]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Type", "Sub item", "Total number of shifts", "Total shift hours", "First name", "Last name",
               "Start Date", "In", "Start - location", "End Date", "Out", "End - location", "Shift hours", None])
    for header, total, shifts in blocks:
        ws.append([header, None, None, total] + [""] * 9)
        for s, e, h in shifts:
            ws.append(["", None, None, "", "A", "B", s, "09:00", None, e, "15:00", None, h])
    if total_row:
        ws["M%d" % (ws.max_row + 1)] = "=SUM(M2:M%d)" % ws.max_row
    wb.save(path)


def make_chart(path, rows):
    """rows: [(family, client, name, ahs, hm, pc, rs)]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ct.DEFAULT_CHART_SHEET
    ws["A2"] = "Code"
    ws.append([])
    ws.cell(3, 1, "Family #"), ws.cell(3, 2, "Client #"), ws.cell(3, 4, "Name")
    ws.cell(3, 6, "AHS ID #"), ws.cell(3, 10, "Homemaking hours")
    ws.cell(3, 11, "Personal Care hours"), ws.cell(3, 12, "Respite hours")
    r = 4
    for fam, cli, name, ahs, hm, pc, rs in rows:
        for c, v in ((1, fam), (2, cli), (4, name), (6, ahs), (10, hm), (11, pc), (12, rs)):
            ws.cell(r, c, v)
        r += 4                        # real chart leaves 3 blank rows under each client
    wb.save(path)


@pytest.fixture
def files(tmp_path):
    tc, ch = tmp_path / "tc.xlsx", tmp_path / "chart.xlsx"
    make_timeclock(tc, [
        ("5  Bob Smith", 14, [(D(2026, 9, 17), D(2026, 9, 17), 6), (D(2026, 9, 20), D(2026, 9, 21), 8)]),
        ("2 Al Jones", 4, [(D(2026, 9, 18), D(2026, 9, 18), 4)]),
        ("99 New Person", 3, [(D(2026, 9, 19), D(2026, 9, 19), 3)]),
    ])
    make_chart(ch, [(2, 3, "Al", "A1", 1, 2, 3), (5, 6, "Bob", "B1", 0, 5, 1), (5, 7, "Bobs wife", "B2", 0, 5, 1)])
    return tc, ch, tmp_path / "out.xlsx"


def run(files, **kw):
    tc, ch, out = files
    args = dict(chart_sheet=ct.DEFAULT_CHART_SHEET, out=out, month=None, rate=34.03, date_field="start")
    args.update(kw)
    return ct.run(tc, ch, args.pop("chart_sheet"), args.pop("out"), args.pop("month"),
                  args.pop("rate"), args.pop("date_field"))


def sheet(out):
    return openpyxl.load_workbook(out)


def test_total_row_is_not_counted(files):
    _, stats = run(files)
    assert stats["hours_in"] == 21 and stats["hours_out"] == 0     # 14 + 4 + 3, not + 21 again


def test_blocks_sorted_and_hours_on_start_date(files):
    out, _ = run(files)
    ws = sheet(out).worksheets[0]
    fams = [ws.cell(ct.FIRST_BLOCK_ROW + i * ct.BLOCK_PITCH, 1).value for i in range(5)]
    assert fams == [2, 5, 5, 99, None]                  # sorted by family; family 5 has two clients
    first5 = ct.FIRST_BLOCK_ROW + ct.BLOCK_PITCH
    pc = first5 + 6
    assert ws.cell(pc, ct.FIRST_DAY_COL + 16).value == 6            # 17th
    assert ws.cell(pc, ct.FIRST_DAY_COL + 19).value == 8            # overnight shift stays on the 20th
    assert ws.cell(first5 + ct.BLOCK_PITCH + 6, ct.FIRST_DAY_COL + 16).value is None   # 2nd client empty


def test_end_date_option_moves_overnight_shift(files):
    out, _ = run(files, date_field="end")
    ws = sheet(out).worksheets[0]
    pc = ct.FIRST_BLOCK_ROW + ct.BLOCK_PITCH + 6
    assert ws.cell(pc, ct.FIRST_DAY_COL + 20).value == 8            # 21st


def test_family_missing_from_chart_is_flagged_not_dropped(files):
    out, stats = run(files)
    assert stats["errors"] == 1
    wb = sheet(out)
    msgs = [r[2] for r in wb["Issues"].iter_rows(min_row=2, values_only=True)]
    assert any("99" in m and "not in the chart" in m for m in msgs)
    ws = wb.worksheets[0]
    top = ct.FIRST_BLOCK_ROW + 3 * ct.BLOCK_PITCH                   # family 99 block
    assert ws.cell(top, 1).value == 99 and ws.cell(top + 6, ct.FIRST_DAY_COL + 18).value == 3


def test_shifts_outside_month_are_reported(tmp_path):
    tc, ch = tmp_path / "tc.xlsx", tmp_path / "chart.xlsx"
    make_timeclock(tc, [("2 Al", 10, [(D(2026, 9, 30), D(2026, 10, 1), 6), (D(2026, 10, 1), D(2026, 10, 1), 4)])])
    make_chart(ch, [(2, 3, "Al", "A1", 1, 2, 3)])
    _, stats = ct.run(tc, ch, ct.DEFAULT_CHART_SHEET, tmp_path / "o.xlsx", "2026-09", 34.03, "start")
    assert stats["hours_in"] == 6 and stats["hours_out"] == 4 and stats["warnings"] >= 1


def test_same_family_twice_is_combined(tmp_path):
    tc, ch = tmp_path / "tc.xlsx", tmp_path / "chart.xlsx"
    make_timeclock(tc, [("2 Al", 2, [(D(2026, 9, 1), D(2026, 9, 1), 2)]),
                        ("2 Al again", 3, [(D(2026, 9, 1), D(2026, 9, 1), 3)])])
    make_chart(ch, [(2, 3, "Al", "A1", 1, 2, 3)])
    out = tmp_path / "o.xlsx"
    ct.run(tc, ch, ct.DEFAULT_CHART_SHEET, out, None, 34.03, "start")
    ws = sheet(out).worksheets[0]
    assert ws.cell(ct.FIRST_BLOCK_ROW + 6, ct.FIRST_DAY_COL).value == 5


def test_header_without_family_number_is_an_error(tmp_path):
    tc, ch = tmp_path / "tc.xlsx", tmp_path / "chart.xlsx"
    make_timeclock(tc, [("Mystery Person", 2, [(D(2026, 9, 1), D(2026, 9, 1), 2)])])
    make_chart(ch, [(2, 3, "Al", "A1", 1, 2, 3)])
    _, stats = ct.run(tc, ch, ct.DEFAULT_CHART_SHEET, tmp_path / "o.xlsx", "2026-09", 34.03, "start")
    assert stats["errors"] == 1 and stats["hours_out"] == 2


def test_wrong_file_gives_readable_error(tmp_path):
    bad = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Qty"])
    wb.save(bad)
    with pytest.raises(ct.InputError, match="Shift hours"):
        ct.read_timeclock(bad)


def test_half_month_dollar_sums(files):
    run(files, month="2026-09")                                     # September: 30 days
    ws = sheet(files[2])["Claim 2026-09"]
    top = ct.FIRST_BLOCK_ROW
    h = top + ct.HALF_OFF
    mid, end = ct.FIRST_DAY_COL + 14, ct.FIRST_DAY_COL + 29        # S = day 15, AH = day 30
    assert ws.cell(h, mid - 2).value == "Days 1-15" and ws.cell(h, end - 2).value == "Days 16-30"
    assert ws.cell(h, mid - 1).value == "Hours" and ws.cell(h, mid).value == "$$"
    assert ws.cell(h + 1, mid).value == f"=SUM(E{top + 14}:S{top + 14})"
    assert ws.cell(h + 2, end).value == f"=SUM(T{top + 15}:AH{top + 15})"
    assert ws.cell(h + 4, end).value == f"=SUM(AH{h + 1}:AH{h + 3})"
    assert ws.cell(h + 1, mid - 1).value == f"=SUM(E{top + 5}:S{top + 5})"
    assert ws.cell(h + 2, end - 1).value == f"=SUM(T{top + 6}:AH{top + 6})"
    assert ws.cell(h + 4, end - 1).value == f"=SUM(AG{h + 1}:AG{h + 3})"
    assert ws.cell(h + 3, mid - 2).value == "Subtotal Respite"
    assert ws.cell(h + 4, end - 2).value == "Subtotal Total"
    assert ws.cell(h + 3, 4).value is None and ct.BLOCK_PITCH > ct.HALF_OFF + 4


def test_period_subtotals_and_client_month_sums(files):
    run(files, month="2026-09")
    ws = sheet(files[2])["Claim 2026-09"]
    top = ct.FIRST_BLOCK_ROW
    h = top + ct.PERIOD_OFF
    assert h < top + ct.HALF_OFF                                    # periods sit above the half months
    spans = [(1, 10), (11, 15), (16, 25), (26, 30)]
    dests = [ct.FIRST_DAY_COL + d2 - 1 for _, d2 in spans]
    for (d1, d2), dest in zip(spans, dests):
        assert ws.cell(h, dest - 2).value == f"Days {d1}-{d2}"
        c1, c2 = ct.col(ct.FIRST_DAY_COL + d1 - 1), ct.col(dest)
        assert ws.cell(h + 1, dest - 1).value == f"=SUM({c1}{top + 5}:{c2}{top + 5})"
        assert ws.cell(h + 3, dest).value == f"=SUM({c1}{top + 16}:{c2}{top + 16})"
    assert ws.cell(h + 2, ct.SUM_COL).value == "Month Personal Care"
    exp = "=" + "+".join(f"{ct.col(d - 1)}{h + 2}" for d in dests)
    assert ws.cell(h + 2, ct.GRAND_COL).value == exp
    exp = "=" + "+".join(f"{ct.col(d)}{h + 4}" for d in dests)
    assert ws.cell(h + 4, ct.GRAND_COL + 1).value == exp
    hh = top + ct.HALF_OFF                                          # half months get month sums too
    assert ws.cell(hh + 1, ct.GRAND_COL).value.count("+") == 1


def test_claimed_table_left_of_half_month_subtotals(files):
    run(files, month="2026-09")
    ws = sheet(files[2])["Claim 2026-09"]
    top = ct.FIRST_BLOCK_ROW
    h = top + ct.HALF_OFF
    for d2 in (15, 30):
        dest = ct.FIRST_DAY_COL + d2 - 1
        lab, dol = ct.claimed_cols(dest)
        assert dol == dest - 5 and lab == dol - 1 and (dest - 2) - dol == 3   # two empty cells before the blue label
        assert ws.cell(h + 1, dest - 4).value is None and ws.cell(h + 1, dest - 3).value is None
        assert ws.cell(h + 1, lab).value == "Claimed Homemaking"
        assert ws.cell(h, dol).value == "$$" and ws.cell(h + 3, dol).value == 0
        assert ws.cell(h + 4, dol).value == f"=SUM({ct.col(dol)}{h + 1}:{ct.col(dol)}{h + 3})"
        assert ws.cell(h + 2, lab).fill.fgColor.rgb[-6:] == ct.PINK
    d1_ = ct.claimed_cols(ct.FIRST_DAY_COL + 14)[1]
    d2_ = ct.claimed_cols(ct.FIRST_DAY_COL + 29)[1]
    assert ws.cell(h + 4, ct.CLAIM_COL + 1).value == f"={ct.col(d1_)}{h + 4}+{ct.col(d2_)}{h + 4}"
    assert ws.cell(top + ct.PERIOD_OFF + 1, ct.CLAIM_COL).value is None   # period table has no claimed table
    last = top + 3 * ct.BLOCK_PITCH
    gh = last + ct.HALF_OFF + 12                                    # all-clients half-month header
    f = ws.cell(gh + 1, d1_).value
    assert f.startswith(f"=SUM({ct.col(d1_)}{h + 1},") and f.count(",") == 3


def test_grand_totals_after_last_block(files):
    run(files, month="2026-09")
    ws = sheet(files[2])["Claim 2026-09"]
    n = 4                                                            # blocks in the fixture
    last = ct.FIRST_BLOCK_ROW + (n - 1) * ct.BLOCK_PITCH
    ph = last + ct.HALF_OFF + 6                                      # 1 blank row under the last block
    hh = ph + 6
    mid, end = ct.FIRST_DAY_COL + 14, ct.FIRST_DAY_COL + 29
    assert ws.cell(ph, ct.FIRST_DAY_COL + 9 - 2).value == "All clients 1-10"
    assert ws.cell(hh, mid - 2).value == "All clients 1-15"
    assert ws.cell(hh + 2, end - 2).fill.fgColor.rgb[-6:] == ct.YELLOW
    f = ws.cell(hh + 4, mid).value
    assert f == f"=SUM(S{hh + 1}:S{hh + 3})"                         # total row adds the three categories
    fh = ws.cell(hh + 2, mid - 1).value
    assert fh.startswith(f"=SUM(R{ct.FIRST_BLOCK_ROW + ct.HALF_OFF + 2},") and fh.count(",") == n - 1
    assert ws.cell(hh + 4, end - 2).value == "Subtotal Total"
    assert ws.cell(hh + 2, ct.GRAND_COL).value == f"=R{hh + 2}+AG{hh + 2}"
    assert ws.cell(hh + 2, ct.GRAND_COL + 1).value == f"=S{hh + 2}+AH{hh + 2}"
    pf = ws.cell(ph + 1, ct.FIRST_DAY_COL + 8).value                 # all clients, 1-10 hours, homemaking
    assert pf.startswith(f"=SUM(M{ct.FIRST_BLOCK_ROW + ct.PERIOD_OFF + 1},") and pf.count(",") == n - 1


def test_dollar_rows_have_month_total(files):
    run(files, month="2026-09")
    ws = sheet(files[2])["Claim 2026-09"]
    top = ct.FIRST_BLOCK_ROW
    for off in (14, 15, 16, 17):
        assert ws.cell(top + off, ct.SUM_COL).value == f"=SUM(E{top + off}:AI{top + off})"


def test_each_block_has_a_border(files):
    run(files, month="2026-09")
    ws = sheet(files[2])["Claim 2026-09"]
    for i in range(4):                                  # includes the two clients of one family
        top = ct.FIRST_BLOCK_ROW + i * ct.BLOCK_PITCH
        assert ws.cell(top, 1).border.top.style == "medium" and ws.cell(top, 1).border.left.style == "medium"
        assert ws.cell(top + ct.HALF_OFF + 4, ct.CLAIM_COL + 1).border.bottom.style == "medium"
        assert ws.cell(top + ct.HALF_OFF + 4, ct.CLAIM_COL + 1).border.right.style == "medium"


def test_subtotal_label_and_ak_highlight_colors(files):
    run(files, month="2026-09")
    ws = sheet(files[2])["Claim 2026-09"]
    top = ct.FIRST_BLOCK_ROW
    rgb = lambda cell: cell.fill.fgColor.rgb[-6:]               # noqa: E731
    for k in range(4):
        for dest in (ct.FIRST_DAY_COL + 14, ct.FIRST_DAY_COL + 29):
            assert rgb(ws.cell(top + ct.HALF_OFF + 1 + k, dest - 2)) == ct.BLUE_M
    assert rgb(ws.cell(top + 7, ct.GRAND_COL)) == ct.YELLOW
    assert ws.cell(top + 7, ct.GRAND_COL).value.startswith("=SUM(")
    assert ws.cell(top + 8, ct.GRAND_COL).fill.fgColor.rgb[-6:] != ct.YELLOW


def test_leading_parentheses_ignored_before_family_number():
    f = ct.strip_leading_parens
    assert f("(回國)67 Rui dad").strip().startswith("67")
    assert f("(回國3個月)59 Lian Mei Zhou").strip().startswith("59")
    assert f("（a）(b) 5 X") .strip().startswith("5")
    assert f("12 Li (Wu)") == "12 Li (Wu)"


def test_office_header_gets_blank_block(tmp_path):
    tc, ch, out = tmp_path / "tc.xlsx", tmp_path / "chart.xlsx", tmp_path / "o.xlsx"
    make_timeclock(tc, [
        ("Office", 5, [(D(2026, 9, 17), D(2026, 9, 17), 5)]),
        ("2 Al Jones", 4, [(D(2026, 9, 18), D(2026, 9, 18), 4)]),
        ("Office Staff", 1, [(D(2026, 9, 19), D(2026, 9, 19), 1)]),       # not a perfect match
    ])
    make_chart(ch, [(2, 3, "Al", "A1", 1, 2, 3)])
    _, stats = ct.run(tc, ch, ct.DEFAULT_CHART_SHEET, out, "2026-09", 34.03, "start")
    ws = openpyxl.load_workbook(out)["Claim 2026-09"]
    top = ct.FIRST_BLOCK_ROW + ct.BLOCK_PITCH                          # Office block is last
    assert [ws.cell(top, c).value for c in (1, 2, 3)] == [None, None, None]
    assert ws.cell(top + 2, 3).value == "Office"
    assert ws.cell(top + 6, ct.FIRST_DAY_COL + 16).value == 5
    assert stats["errors"] == 1                                         # only "Office Staff"


def test_office_match_is_case_insensitive():
    assert all(ct.is_office(h) for h in ("Office", "OFFICE", "office", "  Office "))
    assert not ct.is_office("Office Staff")
